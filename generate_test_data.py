"""
Генератор тестового Excel-файла с учебным планом.
Создаёт Шаблон_импорта.xlsx (template.xlsx) с 10 классами.

Все учителя имеют нагрузку ≤ 40 часов.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


# Классы
CLASSES_5 = ["5А", "5Б", "5В", "5Г", "5Д"]
CLASSES_6 = ["6А", "6Б", "6В", "6Г", "6Д"]

# Предметы (для 5 и 6 классов)
SUBJECTS_5 = [
    ("Русский язык", 5),
    ("Литература", 3),
    ("Иностранный язык", 3),
    ("Математика", 5),
    ("История", 3),
    ("География", 1),
    ("Биология", 1),
    ("Музыка", 1),
    ("Изобразительное искусство", 1),
    ("Труд (технология)", 2),
    ("Физическая культура", 2),
]

SUBJECTS_6 = [
    ("Русский язык", 6),
    ("Литература", 3),
    ("Иностранный язык", 3),
    ("Математика", 5),
    ("История", 3),
    ("География", 1),
    ("Биология", 1),
    ("Музыка", 1),
    ("Изобразительное искусство", 1),
    ("Труд (технология)", 2),
    ("Физическая культура", 2),
]

# Учителя
TEACHERS = [
    ("Иванова А.П.", ["Русский язык", "Литература"]),
    ("Козлова Н.П.", ["Русский язык", "Литература"]),
    ("Смирнова Р.А.", ["Русский язык", "Литература"]),
    ("Петрова Н.С.", ["Русский язык", "Литература"]),
    ("Петров С.И.", ["Математика"]),
    ("Соколов Д.А.", ["Математика"]),
    ("Сидорова Е.В.", ["История", "География"]),
    ("Белов А.И.", ["История"]),
    ("Волков А.В.", ["Биология"]),
    ("Кузнецов П.Р.", ["Физическая культура"]),
    ("Попова М.С.", ["Музыка"]),
    ("Новиков И.Д.", ["Изобразительное искусство"]),
    ("Морозов В.К.", ["Труд (технология)"]),
    ("Орлова Т.С.", ["Иностранный язык"]),
]

# Русский + Литература — 4 учителя делят 10 классов
RU_TEACHERS = {
    "5А": "Иванова А.П.",
    "5Б": "Иванова А.П.",
    "5В": "Козлова Н.П.",
    "5Г": "Козлова Н.П.",
    "5Д": "Смирнова Р.А.",
    "6А": "Смирнова Р.А.",
    "6Б": "Петрова Н.С.",
    "6В": "Петрова Н.С.",
    "6Г": "Иванова А.П.",
    "6Д": "Козлова Н.П.",
}

# Математика — 2 учителя
MATH_TEACHERS = {
    "5А": "Петров С.И.",
    "5Б": "Петров С.И.",
    "5В": "Петров С.И.",
    "5Г": "Петров С.И.",
    "5Д": "Петров С.И.",
    "6А": "Соколов Д.А.",
    "6Б": "Соколов Д.А.",
    "6В": "Соколов Д.А.",
    "6Г": "Соколов Д.А.",
    "6Д": "Соколов Д.А.",
}

# Остальные — один учитель на всё
OTHER_TEACHERS = {
    "Иностранный язык": "Орлова Т.С.",
    "География": "Сидорова Е.В.",
    "Биология": "Волков А.В.",
    "Физическая культура": "Кузнецов П.Р.",
    "Музыка": "Попова М.С.",
    "Изобразительное искусство": "Новиков И.Д.",
    "Труд (технология)": "Морозов В.К.",
}

# Кабинеты
ROOMS = [
    ("101", 30, ""),
    ("102", 30, ""),
    ("103", 30, ""),
    ("104", 30, ""),
    ("105", 30, ""),
    ("106", 30, ""),
    ("107", 30, ""),
    ("Спортзал 1", 40, "Физическая культура"),
    ("Спортзал 2", 40, "Физическая культура"),
    ("Мастерская 1", 25, "Труд (технология)"),
    ("Мастерская 2", 25, "Труд (технология)"),
    ("Кабинет музыки 1", 30, "Музыка"),
    ("Кабинет музыки 2", 30, "Музыка"),
    ("Кабинет ИЗО 1", 30, "Изобразительное искусство"),
    ("Кабинет ИЗО 2", 30, "Изобразительное искусство"),
]


def make_header(ws, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")


def create_school_sheet(wb):
    ws = wb.create_sheet("Школа")
    make_header(ws, ["Название", "Адрес", "Директор", "Телефон", "Email", "Учебный год"])
    ws.cell(row=2, column=1, value="МБОУ СОШ №1")
    ws.cell(row=2, column=2, value="ул. Ленина, 1")
    ws.cell(row=2, column=3, value="Иванов И.И.")
    ws.cell(row=2, column=4, value="+7-999-123-45-67")
    ws.cell(row=2, column=5, value="school@mail.ru")
    ws.cell(row=2, column=6, value="2025-2026")
    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 25


def create_teachers_sheet(wb):
    ws = wb.create_sheet("Учителя")
    make_header(ws, ["ФИО", "Телефон", "Email", "Заметки", "Предметы"])
    for row, (name, subjects) in enumerate(TEACHERS, 2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=5, value=", ".join(subjects))
    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 25


def create_subjects_sheet(wb):
    ws = wb.create_sheet("Предметы")
    make_header(ws, ["Название", "Сложность"])
    all_subjects = sorted(set(s[0] for s in SUBJECTS_5))
    for row, name in enumerate(all_subjects, 2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=5)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12


def create_classes_sheet(wb):
    ws = wb.create_sheet("Классы")
    make_header(ws, ["Название", "Учеников", "Смена"])
    row = 2
    for c in CLASSES_5 + CLASSES_6:
        ws.cell(row=row, column=1, value=c)
        ws.cell(row=row, column=2, value=25)
        ws.cell(row=row, column=3, value=1)
        row += 1
    for col in range(1, 4):
        ws.column_dimensions[chr(64 + col)].width = 20


def create_rooms_sheet(wb):
    ws = wb.create_sheet("Кабинеты")
    make_header(ws, ["Номер кабинета", "Спец. предмет", "Мест"])
    for row, (name, cap, spec) in enumerate(ROOMS, 2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=spec)
        ws.cell(row=row, column=3, value=cap)
    for col in range(1, 4):
        ws.column_dimensions[chr(64 + col)].width = 20


def pick_teacher(class_name, subject):
    if subject in ("Русский язык", "Литература"):
        return RU_TEACHERS.get(class_name, "")
    if subject == "Математика":
        return MATH_TEACHERS.get(class_name, "")
    if subject == "История":
        if class_name.startswith("5"):
            return "Сидорова Е.В."
        return "Белов А.И."
    return OTHER_TEACHERS.get(subject, "")


def create_plan_sheet(wb):
    ws = wb.create_sheet("Учебный план")
    make_header(ws, ["Класс", "Предмет", "Часов", "Учитель"])

    row = 2
    for c in CLASSES_5:
        for subject, hours in SUBJECTS_5:
            ws.cell(row=row, column=1, value=c)
            ws.cell(row=row, column=2, value=subject)
            ws.cell(row=row, column=3, value=hours)
            ws.cell(row=row, column=4, value=pick_teacher(c, subject))
            row += 1

    for c in CLASSES_6:
        for subject, hours in SUBJECTS_6:
            ws.cell(row=row, column=1, value=c)
            ws.cell(row=row, column=2, value=subject)
            ws.cell(row=row, column=3, value=hours)
            ws.cell(row=row, column=4, value=pick_teacher(c, subject))
            row += 1

    for col in range(1, 5):
        ws.column_dimensions[chr(64 + col)].width = 25


def main():
    wb = Workbook()
    wb.remove(wb.active)

    create_school_sheet(wb)
    create_teachers_sheet(wb)
    create_subjects_sheet(wb)
    create_classes_sheet(wb)
    create_rooms_sheet(wb)
    create_plan_sheet(wb)

    wb.save("template.xlsx")

    print("✅ Файл template.xlsx создан")
    print(f"   Классов: {len(CLASSES_5) + len(CLASSES_6)}")
    print(f"   Учителей: {len(TEACHERS)}")
    print(f"   Предметов: {len(set(s[0] for s in SUBJECTS_5))}")
    print(f"   Кабинетов: {len(ROOMS)}")


if __name__ == "__main__":
    main()